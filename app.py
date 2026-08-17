from __future__ import annotations

import io
import json
import os
import re
import secrets
from datetime import datetime, timedelta, timezone
from functools import wraps
from pathlib import Path
from typing import Any, Callable, TypeVar, cast

from flask import (
    Flask,
    current_app,
    flash,
    g,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import ValidationError
from werkzeug.security import check_password_hash, generate_password_hash

from src.database import get_db, init_app as init_database
from src.shelf_life_chain import POLYNOMIAL_TEMPERATURE_RANGE_C, build_shelf_life_chain


PREDICTION_METHODS = ["polynomial_regression", "llm_structured"]
FIXED_TEMPERATURES_C = [0, 5, 10, 15, 20, 25]
USERNAME_PATTERN = re.compile(r"^[\w\u4e00-\u9fff-]+$", re.UNICODE)
F = TypeVar("F", bound=Callable[..., Any])


def _load_or_create_secret(instance_path: str) -> str:
    configured_secret = os.getenv("FLASK_SECRET_KEY")
    if configured_secret:
        return configured_secret

    secret_path = Path(instance_path) / ".secret_key"
    secret_path.parent.mkdir(parents=True, exist_ok=True)
    if secret_path.exists():
        return secret_path.read_text(encoding="utf-8").strip()

    secret = secrets.token_hex(32)
    secret_path.write_text(secret, encoding="utf-8")
    return secret


def create_app(test_config: dict[str, Any] | None = None) -> Flask:
    app = Flask(__name__, instance_relative_config=True)
    app.config.from_mapping(
        SECRET_KEY=_load_or_create_secret(app.instance_path),
        DATABASE=str(Path(app.instance_path) / "meizao.sqlite3"),
        PERMANENT_SESSION_LIFETIME=timedelta(days=7),
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
        SESSION_COOKIE_SECURE=False,
        CSRF_ENABLED=True,
    )
    if test_config:
        app.config.update(test_config)

    init_database(app)
    shelf_life_chain = build_shelf_life_chain()

    @app.before_request
    def load_logged_in_user() -> None:
        user_id = session.get("user_id")
        g.user = None
        if user_id is not None:
            g.user = get_db().execute(
                "SELECT id, username, created_at FROM users WHERE id = ?", (user_id,)
            ).fetchone()
            if g.user is None:
                session.clear()

    @app.context_processor
    def inject_template_helpers() -> dict[str, Any]:
        return {"csrf_token": _csrf_token}

    def login_required(view: F) -> F:
        @wraps(view)
        def wrapped_view(*args: Any, **kwargs: Any):
            if g.user is None:
                if request.is_json or request.path.startswith("/api/"):
                    return jsonify({"error": "authentication_required", "message": "请先登录。"}), 401
                return redirect(url_for("login", next=request.full_path if request.query_string else request.path))
            return view(*args, **kwargs)

        return cast(F, wrapped_view)

    @app.route("/register", methods=["GET", "POST"])
    def register():
        if g.user is not None:
            return redirect(url_for("index"))

        if request.method == "POST":
            if not _valid_csrf():
                flash("页面已过期，请重新提交。", "error")
                return render_template("auth.html", mode="register"), 400

            username = request.form.get("username", "").strip()
            password = request.form.get("password", "")
            confirm_password = request.form.get("confirm_password", "")
            error = _validate_registration(username, password, confirm_password)
            database = get_db()

            if error is None and database.execute(
                "SELECT id FROM users WHERE username = ?", (username,)
            ).fetchone():
                error = "该用户名已被注册。"

            if error is None:
                cursor = database.execute(
                    "INSERT INTO users (username, password_hash) VALUES (?, ?)",
                    (username, generate_password_hash(password)),
                )
                database.commit()
                session.clear()
                session["user_id"] = cursor.lastrowid
                session.permanent = True
                _csrf_token()
                flash("注册成功，欢迎使用货架期预测系统。", "success")
                return redirect(url_for("index"))

            flash(error, "error")

        return render_template("auth.html", mode="register")

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if g.user is not None:
            return redirect(url_for("index"))

        if request.method == "POST":
            if not _valid_csrf():
                flash("页面已过期，请重新提交。", "error")
                return render_template("auth.html", mode="login"), 400

            username = request.form.get("username", "").strip()
            password = request.form.get("password", "")
            user = get_db().execute(
                "SELECT id, username, password_hash FROM users WHERE username = ?", (username,)
            ).fetchone()

            if user is None or not check_password_hash(user["password_hash"], password):
                flash("用户名或密码错误。", "error")
            else:
                session.clear()
                session["user_id"] = user["id"]
                session.permanent = True
                _csrf_token()
                next_url = request.args.get("next", "")
                return redirect(next_url if _is_safe_next(next_url) else url_for("index"))

        return render_template("auth.html", mode="login")

    @app.post("/logout")
    @login_required
    def logout():
        if not _valid_csrf():
            flash("退出请求已过期，请重试。", "error")
            return redirect(url_for("index"))
        session.clear()
        flash("你已安全退出。", "success")
        return redirect(url_for("login"))

    @app.get("/")
    @login_required
    def index():
        history = _get_history(g.user["id"], limit=100)
        return render_template(
            "index.html",
            fixed_temperatures_c=FIXED_TEMPERATURES_C,
            prediction_methods=PREDICTION_METHODS,
            history=history,
            history_summary=_history_summary(history),
        )

    @app.get("/health")
    def health():
        return jsonify({"status": "ok"})

    @app.get("/methods")
    def methods():
        return jsonify(
            {
                "prediction_methods": PREDICTION_METHODS,
                "default_prediction_method": "polynomial_regression",
                "fixed_temperatures_c": FIXED_TEMPERATURES_C,
                "temperature_range_c": POLYNOMIAL_TEMPERATURE_RANGE_C,
                "llm_structured_provider": "siliconflow",
                "llm_structured_requires": ["SILICONFLOW_API_KEY", "langchain-openai"],
            }
        )

    @app.post("/predict")
    @login_required
    def predict():
        if not _valid_csrf():
            return jsonify({"error": "csrf_error", "message": "页面已过期，请刷新后重试。"}), 400

        payload = request.get_json(silent=True) or {}
        try:
            result = shelf_life_chain.invoke(payload)
        except ValidationError as exc:
            return jsonify({"error": "validation_error", "details": exc.errors()}), 400
        except ValueError as exc:
            return jsonify({"error": "invalid_input", "message": str(exc)}), 400
        except Exception as exc:
            return (
                jsonify(
                    {
                        "error": "prediction_failed",
                        "message": str(exc),
                        "hint": "For llm_structured, check SiliconFlow network access, API key, model name, and proxy settings.",
                    }
                ),
                502,
            )

        history_id, created_at = _save_prediction(g.user["id"], result)
        response = dict(result)
        response["history_id"] = history_id
        response["created_at"] = created_at
        return jsonify(response), 200

    @app.get("/api/history")
    @login_required
    def history_api():
        limit = min(max(request.args.get("limit", 100, type=int), 1), 500)
        history = _get_history(g.user["id"], limit=limit)
        return jsonify({"records": history, "summary": _history_summary(history)})

    @app.post("/history/clear")
    @login_required
    def clear_history():
        if not _valid_csrf():
            if request.is_json:
                return jsonify({"error": "csrf_error", "message": "页面已过期，请刷新后重试。"}), 400
            flash("清除请求已过期，请重试。", "error")
            return redirect(url_for("index"))

        cursor = get_db().execute(
            "DELETE FROM prediction_history WHERE user_id = ?", (g.user["id"],)
        )
        get_db().commit()
        flash(f"已清除 {cursor.rowcount} 条预测历史。", "success")
        return redirect(url_for("index"))

    @app.get("/history/export")
    @login_required
    def export_history():
        history = _get_history(g.user["id"], limit=None)
        workbook = _build_history_workbook(history, g.user["username"])
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        date_text = datetime.now().strftime("%Y%m%d")
        return send_file(
            output,
            as_attachment=True,
            download_name=f"樱桃预测历史_{date_text}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    return app


def _csrf_token() -> str:
    token = session.get("_csrf_token")
    if not token:
        token = secrets.token_urlsafe(32)
        session["_csrf_token"] = token
    return token


def _valid_csrf() -> bool:
    if not current_app.config.get("CSRF_ENABLED", True):
        return True
    expected = session.get("_csrf_token", "")
    supplied = request.headers.get("X-CSRF-Token") or request.form.get("csrf_token", "")
    return bool(expected and supplied and secrets.compare_digest(expected, supplied))


def _validate_registration(username: str, password: str, confirmation: str) -> str | None:
    if not 3 <= len(username) <= 30:
        return "用户名长度应为 3-30 个字符。"
    if not USERNAME_PATTERN.fullmatch(username):
        return "用户名只能包含中文、字母、数字、下划线或连字符。"
    if len(password) < 8:
        return "密码至少需要 8 个字符。"
    if password != confirmation:
        return "两次输入的密码不一致。"
    return None


def _is_safe_next(next_url: str) -> bool:
    return bool(next_url and next_url.startswith("/") and not next_url.startswith("//"))


def _save_prediction(user_id: int, result: dict[str, Any]) -> tuple[int, str]:
    created_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    shelf_range = result["shelf_life_range_days"]
    cursor = get_db().execute(
        """
        INSERT INTO prediction_history (
            user_id, storage_temperature_c, handheld_hardness, firmness_g_mm2,
            estimated_shelf_life_days, range_min_days, range_max_days, risk_level,
            confidence, prediction_method, decision, recommendations_json, created_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            user_id,
            result["storage_temperature_c"],
            result["handheld_hardness"],
            result["firmness_g_mm2"],
            result["estimated_shelf_life_days"],
            shelf_range["min"],
            shelf_range["max"],
            result["risk_level"],
            result["confidence"],
            result["prediction_method"],
            result["decision"],
            json.dumps(result["recommendations"], ensure_ascii=False),
            created_at,
        ),
    )
    get_db().commit()
    return int(cursor.lastrowid), created_at


def _get_history(user_id: int, limit: int | None) -> list[dict[str, Any]]:
    sql = """
        SELECT id, storage_temperature_c, handheld_hardness, firmness_g_mm2,
               estimated_shelf_life_days, range_min_days, range_max_days,
               risk_level, confidence, prediction_method, decision,
               recommendations_json, created_at
        FROM prediction_history
        WHERE user_id = ?
        ORDER BY created_at DESC, id DESC
    """
    params: list[Any] = [user_id]
    if limit is not None:
        sql += " LIMIT ?"
        params.append(limit)

    rows = get_db().execute(sql, params).fetchall()
    records: list[dict[str, Any]] = []
    for row in rows:
        item = dict(row)
        item["recommendations"] = json.loads(item.pop("recommendations_json") or "[]")
        records.append(item)
    return records


def _history_summary(history: list[dict[str, Any]]) -> dict[str, Any]:
    total = len(history)
    risk_counts = {"low": 0, "medium": 0, "high": 0}
    for item in history:
        risk = item["risk_level"]
        if risk in risk_counts:
            risk_counts[risk] += 1
    average_days = round(
        sum(float(item["estimated_shelf_life_days"]) for item in history) / total, 1
    ) if total else 0
    return {"total": total, "average_days": average_days, "risk_counts": risk_counts}


def _build_history_workbook(history: list[dict[str, Any]], username: str) -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "预测历史"
    headers = [
        "序号", "预测时间（UTC）", "用户", "贮藏温度（℃）", "手持硬度",
        "映射硬度（g·mm⁻²）", "预计货架期（天）", "区间下限（天）",
        "区间上限（天）", "风险等级", "置信度", "预测方法", "处理建议",
    ]
    worksheet.append(headers)
    risk_labels = {"low": "低", "medium": "中", "high": "高"}
    method_labels = {"polynomial_regression": "多项式回归", "llm_structured": "LLM 结构化判断"}

    for index, item in enumerate(history, start=1):
        worksheet.append(
            [
                index,
                item["created_at"],
                username,
                item["storage_temperature_c"],
                item["handheld_hardness"],
                item["firmness_g_mm2"],
                item["estimated_shelf_life_days"],
                item["range_min_days"],
                item["range_max_days"],
                risk_labels.get(item["risk_level"], item["risk_level"]),
                item["confidence"],
                method_labels.get(item["prediction_method"], item["prediction_method"]),
                item["decision"],
            ]
        )

    header_fill = PatternFill("solid", fgColor="0F766E")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    widths = [8, 24, 16, 16, 14, 22, 18, 16, 16, 12, 12, 20, 42]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return workbook


app = create_app()


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=True)


