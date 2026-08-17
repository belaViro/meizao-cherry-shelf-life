from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path.cwd()
OUT_DIR = ROOT / '软著材料'
IMG_DIR = OUT_DIR / '用户说明书图片'
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_FILE = OUT_DIR / '樱桃货架期预测与数据分析系统V1.0用户操作说明书.docx'

SOFTWARE = '樱桃货架期预测与数据分析系统'
VERSION = 'V1.0'
DOC_TITLE = f'{SOFTWARE} {VERSION} 用户操作说明书'
BLACK = RGBColor(0, 0, 0)


def set_east_asia_font(run, font_name):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.color.rgb = BLACK


def set_cell_shading(cell, fill='F2F2F2'):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for margin, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{margin}'))
        if node is None:
            node = OxmlElement(f'w:{margin}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def add_page_field(paragraph, prefix='', suffix=''):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if prefix:
        run = paragraph.add_run(prefix)
        set_east_asia_font(run, '宋体')
        run.font.size = Pt(9)
    run = paragraph.add_run()
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = ' PAGE '
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_east_asia_font(run, '宋体')
    run.font.size = Pt(9)
    if suffix:
        run = paragraph.add_run(suffix)
        set_east_asia_font(run, '宋体')
        run.font.size = Pt(9)


def set_page_number_start(section, start=1):
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn('w:pgNumType'))
    if pg_num_type is None:
        pg_num_type = OxmlElement('w:pgNumType')
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn('w:start'), str(start))


def add_toc(document):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run()
    fld_char = OxmlElement('w:fldChar')
    fld_char.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' TOC \\o "1-3" \\z \\u '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    placeholder = OxmlElement('w:t')
    placeholder.text = '目录将在打开文档时自动更新。'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([fld_char, instr, separate, placeholder, end])
    set_east_asia_font(run, '宋体')
    run.font.size = Pt(10.5)


def configure_document(document):
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)

    normal = document.styles['Normal']
    normal.font.name = '宋体'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.widow_control = True

    title = document.styles['Title']
    title.font.name = '黑体'
    title._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = BLACK

    heading_specs = {
        'Heading 1': ('黑体', 16, True, 12, 8),
        'Heading 2': ('黑体', 14, True, 8, 4),
        'Heading 3': ('宋体', 10.5, True, 5, 2),
    }
    for style_name, (font_name, size, bold, before, after) in heading_specs.items():
        style = document.styles[style_name]
        style.font.name = font_name
        style._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)

    if '图题' not in document.styles:
        caption = document.styles.add_style('图题', WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = document.styles['图题']
    caption.font.name = '宋体'
    caption._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    caption.font.size = Pt(9)
    caption.font.color.rgb = BLACK
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.first_line_indent = Cm(0)

    for toc_name in ('TOC 1', 'TOC 2', 'TOC 3'):
        if toc_name in document.styles:
            toc_style = document.styles[toc_name]
            toc_style.font.name = '宋体'
            toc_style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            toc_style.font.color.rgb = BLACK
            toc_style.font.size = Pt(10.5)


def add_heading(document, text, level=1):
    p = document.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_para(document, text='', bold_prefix=None, center=False, indent=True):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.74) if indent and not center else Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        set_east_asia_font(r1, '宋体')
        r1.font.size = Pt(10.5)
        r2 = p.add_run(text[len(bold_prefix):])
        set_east_asia_font(r2, '宋体')
        r2.font.size = Pt(10.5)
    else:
        r = p.add_run(text)
        set_east_asia_font(r, '宋体')
        r.font.size = Pt(10.5)
    return p


def add_steps(document, steps):
    for i, text in enumerate(steps, 1):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.74)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        r1 = p.add_run(f'（{i}）')
        r1.bold = True
        set_east_asia_font(r1, '宋体')
        r1.font.size = Pt(10.5)
        r2 = p.add_run(text)
        set_east_asia_font(r2, '宋体')
        r2.font.size = Pt(10.5)


def add_bullets(document, items):
    for text in items:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        r = p.add_run('● ' + text)
        set_east_asia_font(r, '宋体')
        r.font.size = Pt(10.5)


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False if widths else True
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for idx, header in enumerate(headers):
        cell = header_row.cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, 'E7E6E6')
        set_cell_margins(cell)
        if widths:
            cell.width = Cm(widths[idx])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(str(header))
        r.bold = True
        set_east_asia_font(r, '宋体')
        r.font.size = Pt(9)
    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Cm(widths[idx])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(value))
            set_east_asia_font(r, '宋体')
            r.font.size = Pt(9)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_figure(document, filename, caption, max_width_cm=14.6, max_height_cm=17.5):
    path = IMG_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    with Image.open(path) as image:
        width_px, height_px = image.size
    ratio = width_px / height_px
    width_cm = max_width_cm
    height_cm = width_cm / ratio
    if height_cm > max_height_cm:
        height_cm = max_height_cm
        width_cm = height_cm * ratio
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm), height=Cm(height_cm))
    cp = document.add_paragraph(style='图题')
    cr = cp.add_run(caption)
    set_east_asia_font(cr, '宋体')
    cr.font.size = Pt(9)


def new_chapter(document, title):
    document.add_page_break()
    add_heading(document, title, 1)


doc = Document()
configure_document(doc)
doc.core_properties.title = DOC_TITLE
doc.core_properties.subject = '计算机软件著作权登记用户操作说明书'
doc.core_properties.author = SOFTWARE
doc.core_properties.keywords = '软件著作权,用户操作说明书,樱桃货架期预测,数据分析'
doc.core_properties.comments = ''

# Cover
for _ in range(5):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Cm(0)
r = p.add_run(SOFTWARE)
set_east_asia_font(r, '黑体')
r.font.size = Pt(24)
r.font.bold = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Cm(0)
r = p.add_run(VERSION)
set_east_asia_font(r, '黑体')
r.font.size = Pt(20)
r.font.bold = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Cm(0)
p.paragraph_format.space_before = Pt(24)
r = p.add_run('用户操作说明书')
set_east_asia_font(r, '黑体')
r.font.size = Pt(22)
r.font.bold = True
for _ in range(8):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Cm(0)
r = p.add_run('二〇二六年八月')
set_east_asia_font(r, '宋体')
r.font.size = Pt(14)

# Main section
main = doc.add_section(WD_SECTION.NEW_PAGE)
main.page_width = Cm(21)
main.page_height = Cm(29.7)
main.top_margin = Cm(2.54)
main.bottom_margin = Cm(2.54)
main.left_margin = Cm(3.17)
main.right_margin = Cm(3.17)
main.header_distance = Cm(1.5)
main.footer_distance = Cm(1.5)
main.header.is_linked_to_previous = False
main.footer.is_linked_to_previous = False
set_page_number_start(main, 1)
header_p = main.header.paragraphs[0]
header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
header_p.paragraph_format.first_line_indent = Cm(0)
hr = header_p.add_run(DOC_TITLE)
set_east_asia_font(hr, '宋体')
hr.font.size = Pt(9)
footer_p = main.footer.paragraphs[0]
add_page_field(footer_p, '第 ', ' 页')

# Contents
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Cm(0)
r = p.add_run('目  录')
set_east_asia_font(r, '黑体')
r.font.size = Pt(18)
r.font.bold = True
p.paragraph_format.space_after = Pt(18)
add_toc(doc)

# 1
new_chapter(doc, '1 引言')
add_heading(doc, '1.1 编写目的', 2)
add_para(doc, '本说明书用于说明樱桃货架期预测与数据分析系统V1.0的运行环境、功能组成和操作方法，使使用人员能够完成用户注册、系统登录、货架期预测、结果查看、数据分析、历史记录管理和Excel数据导出等操作。')
add_para(doc, '本说明书适用于软件安装、使用、维护和功能验收，也可作为软件使用培训和运行管理的依据。')
add_heading(doc, '1.2 软件背景', 2)
add_para(doc, '樱桃在贮藏和流通过程中，其品质变化与贮藏温度、果实硬度等条件密切相关。本系统以贮藏温度和手持式硬度计读数为主要输入，通过本地预测模型计算预计货架期、预测区间和风险等级，并自动形成个人预测历史。')
add_para(doc, '系统采用浏览器操作方式，将预测、统计、图表和历史数据管理集中在统一工作台中，便于对不同贮藏条件下的预测结果进行记录、比较和导出。')
add_heading(doc, '1.3 软件名称及版本', 2)
add_table(doc, ['项目', '内容'], [
    ['软件全称', SOFTWARE],
    ['软件简称', '樱桃货架期预测系统'],
    ['版本号', VERSION],
    ['软件类型', '基于浏览器的应用软件'],
    ['主要用途', '樱桃货架期预测、风险判断、数据分析和历史记录管理'],
], [3.4, 11.0])
add_heading(doc, '1.4 适用对象', 2)
add_bullets(doc, [
    '需要记录樱桃贮藏条件并获得货架期参考结果的使用人员。',
    '需要对预测历史进行统计、比较和导出的数据管理人员。',
    '负责软件部署、运行维护和账号管理的技术人员。',
])
add_heading(doc, '1.5 术语和缩略语', 2)
add_table(doc, ['术语', '说明'], [
    ['货架期', '在给定贮藏条件下，产品保持可接受品质的预计时间。'],
    ['手持硬度', '使用手持式硬度计获得的果实硬度读数，系统允许输入50至90。'],
    ['映射硬度', '由手持硬度通过分段线性换算得到的g·mm⁻²数值。'],
    ['预测区间', '围绕预计货架期给出的参考上下限范围。'],
    ['风险等级', '根据温度、硬度和预测天数划分的低、中、高风险。'],
    ['置信度', '系统根据输入条件给出的预测可靠程度参考值。'],
    ['CSRF', '用于校验页面提交来源、降低跨站请求伪造风险的安全机制。'],
], [3.4, 11.0])

# 2
new_chapter(doc, '2 软件概述')
add_heading(doc, '2.1 软件简介', 2)
add_para(doc, '樱桃货架期预测与数据分析系统V1.0采用Flask构建服务端应用，以SQLite保存账号和预测历史，通过HTML、CSS和JavaScript提供浏览器操作界面。系统支持用户独立账号、预测参数录入、结果计算、趋势图展示、风险分布统计、历史列表、Excel导出和历史清除。')
add_para(doc, '预测结果与当前登录用户关联。不同用户登录后仅能查看、导出和清除本账号产生的历史数据。')
add_heading(doc, '2.2 建设目标', 2)
add_bullets(doc, [
    '统一录入贮藏温度和果实硬度，减少分散记录造成的数据遗漏。',
    '自动计算预计货架期、预测区间、风险等级和处理建议。',
    '自动保存预测历史，形成可追溯的数据记录。',
    '通过趋势图、风险分布和统计指标辅助查看历史变化。',
    '支持导出Excel文件，便于归档、汇总和二次分析。',
])
add_heading(doc, '2.3 软件功能组成', 2)
add_table(doc, ['序号', '功能模块', '功能说明'], [
    ['1', '用户管理', '提供注册、登录、登录状态保持和安全退出功能。'],
    ['2', '预测参数管理', '选择贮藏温度，输入手持硬度并查看映射硬度。'],
    ['3', '货架期预测', '计算预计货架期、预测区间、风险等级和置信度。'],
    ['4', '结果建议', '根据风险等级显示处理建议和操作建议。'],
    ['5', '数据分析', '统计预测次数、平均货架期及各风险等级数量。'],
    ['6', '图表展示', '展示最近20次预测趋势和最近100条风险分布。'],
    ['7', '历史记录', '按照时间倒序显示当前账号最近100条预测记录。'],
    ['8', 'Excel导出', '导出当前账号全部预测历史。'],
    ['9', '历史清除', '经确认后清除当前账号全部预测历史。'],
], [1.3, 3.4, 9.7])
add_heading(doc, '2.4 软件主要特点', 2)
add_para(doc, '系统采用本地多项式回归模型与硬度修正模型进行计算，并使用连续平滑校准函数匹配已知参考点。输入温度与硬度发生变化时，预测结果随参数连续变化，不采用针对单个输入组合直接返回固定结果的处理方式。')
add_para(doc, '系统将预测操作和数据分析集中在同一页面，预测完成后自动刷新统计指标、趋势图、风险分布和历史列表，无需重复切换页面。')
add_heading(doc, '2.5 软件操作流程', 2)
add_steps(doc, [
    '首次使用时进入注册页面，创建个人账号。',
    '使用注册成功的用户名和密码登录系统。',
    '在预测工作台选择贮藏温度并输入手持硬度。',
    '确认预测方法后提交预测任务。',
    '查看预计货架期、预测区间、风险等级和处理建议。',
    '在数据分析区域查看统计指标、趋势图和历史记录。',
    '根据需要导出Excel文件或清除当前账号的历史数据。',
    '操作完成后退出系统。',
])
add_figure(doc, '图2-1_系统主界面全景.png', '图2-1 系统主界面全景', max_height_cm=18.0)

# 3
new_chapter(doc, '3 运行环境')
add_heading(doc, '3.1 硬件环境', 2)
add_table(doc, ['项目', '最低配置', '推荐配置'], [
    ['处理器', '双核处理器', '四核及以上处理器'],
    ['内存', '4 GB', '8 GB及以上'],
    ['可用磁盘空间', '500 MB', '1 GB及以上'],
    ['显示分辨率', '1366×768', '1920×1080及以上'],
    ['网络', '本机运行可不连接外网', '使用远程服务时保持网络稳定'],
], [3.5, 5.4, 5.4])
add_heading(doc, '3.2 软件环境', 2)
add_table(doc, ['类别', '软件或组件'], [
    ['操作系统', 'Windows 10、Windows 11或兼容的服务器操作系统'],
    ['运行环境', 'Python 3.10及以上版本'],
    ['Web框架', 'Flask'],
    ['数据存储', 'SQLite'],
    ['浏览器', 'Microsoft Edge、Google Chrome等现代浏览器'],
    ['数据导出', 'openpyxl，导出格式为.xlsx'],
], [4.0, 10.3])
add_heading(doc, '3.3 浏览器要求', 2)
add_para(doc, '浏览器应启用JavaScript和Cookie。JavaScript用于提交预测请求、更新结果和绘制图表；Cookie用于维持当前用户的登录会话。推荐使用最新稳定版本的Microsoft Edge或Google Chrome，并将页面缩放比例设置为100%。')
add_heading(doc, '3.4 系统启动', 2)
add_para(doc, '进入软件目录，在已安装项目依赖的Python环境中执行以下命令：', indent=False)
add_table(doc, ['操作', '命令'], [
    ['安装依赖', 'python -m pip install -r requirements.txt'],
    ['启动系统', 'python app.py'],
], [3.4, 10.9])
add_para(doc, '服务启动后默认监听本机5000端口。命令行窗口应保持运行状态，关闭该窗口将停止软件服务。')
add_heading(doc, '3.5 系统访问', 2)
add_para(doc, '在浏览器地址栏输入http://127.0.0.1:5000并访问。未登录时系统自动进入登录页面；首次使用人员可通过登录页面下方的注册链接创建账号。')
add_heading(doc, '3.6 系统停止', 2)
add_para(doc, '在运行服务的命令行窗口中按下Ctrl+C组合键，等待服务停止。停止服务前应确认当前预测和导出操作已经完成。')

# 4
new_chapter(doc, '4 用户注册与登录')
add_heading(doc, '4.1 注册与登录页面', 2)
add_para(doc, '系统使用统一账号体系管理用户访问。注册页面用于创建新账号，登录页面用于验证已有账号。账号验证通过后进入预测工作台。')
add_figure(doc, '图4-1_用户登录界面.png', '图4-1 用户登录界面', max_width_cm=9.5, max_height_cm=14.5)
add_heading(doc, '4.2 用户注册', 2)
add_steps(doc, [
    '在登录页面下方选择“注册账号”，进入注册页面。',
    '在“用户名”输入框中输入3至30个字符。用户名可以包含中文、字母、数字、下划线或连字符。',
    '在“密码”输入框中输入不少于8个字符的密码。',
    '在“确认密码”输入框中再次输入相同密码。',
    '选择“创建账号”按钮。注册成功后系统建立用户数据并进入预测工作台。',
])
add_figure(doc, '图4-2_用户注册界面.png', '图4-2 用户注册界面', max_width_cm=9.5, max_height_cm=15.0)
add_heading(doc, '4.3 注册信息校验', 2)
add_table(doc, ['校验项目', '校验规则', '处理结果'], [
    ['用户名长度', '3至30个字符', '不符合时显示用户名长度提示。'],
    ['用户名字符', '中文、字母、数字、下划线或连字符', '包含其他字符时拒绝提交。'],
    ['密码长度', '不少于8个字符', '长度不足时显示密码提示。'],
    ['密码确认', '两次输入必须一致', '不一致时不创建账号。'],
    ['用户名唯一性', '用户名不得与已有账号重复', '重复时提示用户名已存在。'],
], [3.2, 6.0, 5.1])
add_heading(doc, '4.4 用户登录', 2)
add_steps(doc, [
    '进入用户登录页面。',
    '输入已注册的用户名。',
    '输入与该用户名对应的密码。',
    '选择“登录”按钮。',
    '验证成功后进入预测工作台，页面右上角显示当前用户名。',
])
add_heading(doc, '4.5 登录异常处理', 2)
add_para(doc, '用户名不存在或密码不正确时，系统停留在登录页面并显示错误信息。使用人员应检查用户名、密码、输入法状态和字母大小写后重新提交。连续操作时应避免向其他人员透露账号密码。')
add_heading(doc, '4.6 用户退出', 2)
add_para(doc, '选择页面右上角的“退出”按钮，系统清除当前登录会话并返回登录页面。公共计算机或多人共用设备完成操作后应执行退出操作。')

# 5
new_chapter(doc, '5 货架期预测')
add_heading(doc, '5.1 预测工作台', 2)
add_para(doc, '登录成功后进入预测工作台。页面顶部显示软件名称、当前用户名和服务状态，中部左侧为预测参数区域，中部右侧为本次预测结果区域，下部为数据分析与历史记录区域。')
add_figure(doc, '图5-1_预测工作台.png', '图5-1 预测工作台', max_width_cm=14.6, max_height_cm=7.0)
add_heading(doc, '5.2 设置贮藏温度', 2)
add_para(doc, '预测参数区域提供0℃、5℃、10℃、15℃、20℃和25℃六个温度按钮。选择按钮后，该按钮显示选中状态。页面默认选中5℃。每次提交预测时，系统使用当前选中的温度。')
add_heading(doc, '5.3 设置果实硬度', 2)
add_para(doc, '在“手持式硬度计读数”输入框中输入50至90之间的数值。系统根据分段线性换算关系同步显示映射硬度，单位为g·mm⁻²。硬度输入用于修正温度基础货架期，并影响风险等级和预测区间。')
add_table(doc, ['手持硬度', '映射硬度', '说明'], [
    ['50', '200 g·mm⁻²', '允许输入范围下限'],
    ['70', '380 g·mm⁻²', '中间参考点'],
    ['90', '500 g·mm⁻²', '允许输入范围上限'],
], [3.4, 4.5, 6.4])
add_heading(doc, '5.4 选择预测方法', 2)
add_para(doc, '“多项式回归”用于本地快速计算，是当前页面默认选择。预测方法区域会根据运行配置显示其他方法的可用状态；不可用的方法以灰色状态显示且不能选择。')
add_heading(doc, '5.5 提交预测任务', 2)
add_steps(doc, [
    '确认页面显示的当前登录用户正确。',
    '选择需要评估的贮藏温度。',
    '输入实际测得的手持硬度。',
    '确认映射硬度显示正常。',
    '确认预测方法为可用状态。',
    '选择“计算并保存预测”按钮。',
])
add_heading(doc, '5.6 查看预测状态', 2)
add_para(doc, '提交过程中按钮文字显示“计算中”。计算成功后，右侧结果区域显示预测数据，状态更新为已保存；同时统计指标、图表和历史记录自动刷新。预测请求失败时，参数区域显示错误信息，失败数据不会写入历史记录。')
add_heading(doc, '5.7 修改预测参数', 2)
add_para(doc, '完成一次预测后，可直接重新选择温度或修改硬度。修改硬度时，映射硬度同步更新。新的预测不会覆盖原有历史，而是作为新记录保存。')
add_heading(doc, '5.8 重新执行预测', 2)
add_para(doc, '参数修改完成后再次选择“计算并保存预测”。系统按照新的参数重新计算，并将最新记录显示在历史列表顶部。')
add_figure(doc, '图5-2_预测参数与结果区域.png', '图5-2 预测参数与结果区域', max_width_cm=14.6, max_height_cm=11.0)

# 6
new_chapter(doc, '6 预测结果查看')
add_heading(doc, '6.1 预测结果页面', 2)
add_para(doc, '预测完成后，本次预测结果区域显示预计货架期、预测区间、风险等级、处理建议和操作建议。图6-1展示了25℃、手持硬度60条件下的预测结果。')
add_figure(doc, '图6-1_预测结果界面.png', '图6-1 预测结果界面', max_width_cm=14.6, max_height_cm=11.5)
add_heading(doc, '6.2 预计货架期', 2)
add_para(doc, '预计货架期以“天”为单位，是系统根据贮藏温度、手持硬度、映射硬度、基础多项式和校准模型计算得到的中心结果。该结果用于不同贮藏方案之间的比较和管理参考。')
add_heading(doc, '6.3 货架期预测范围', 2)
add_para(doc, '预测区间由下限和上限组成。系统根据温度和硬度条件设置区间宽度，高温或低硬度条件下的区间相对更宽。预测区间用于表示实际环境波动、个体差异等因素可能造成的结果变化。')
add_heading(doc, '6.4 风险等级', 2)
add_table(doc, ['风险等级', '判断特征', '页面处理建议'], [
    ['低风险', '温度较低、硬度较好且预计货架期较长', '适合冷链贮藏，可按正常销售节奏安排。'],
    ['中风险', '温度、硬度或预计天数达到中风险条件', '建议优先销售，并保持贮藏温度稳定。'],
    ['高风险', '温度达到20℃及以上、预计货架期较短或映射硬度偏低', '建议尽快销售或加工，不宜安排长时间贮藏。'],
], [2.6, 5.6, 6.1])
add_heading(doc, '6.5 预测置信度', 2)
add_para(doc, '置信度以百分比形式显示。系统根据温度区间和映射硬度范围计算参考置信度。置信度用于辅助理解预测结果，不替代现场品质检测。')
add_heading(doc, '6.6 决策建议', 2)
add_para(doc, '处理建议与风险等级对应。低风险记录按照正常销售节奏安排；中风险记录优先销售并保持温度稳定；高风险记录尽快销售或加工。操作建议同时提示维持包装完整、减少开箱、定期复测硬度和检查腐烂情况。')
add_heading(doc, '6.7 模型计算信息', 2)
add_para(doc, '本地预测以温度三次多项式作为基础货架期函数，并结合硬度修正系数和温度—硬度平滑校准系数形成最终结果。基础函数如下：', indent=False)
add_para(doc, 'L₀(T) = -0.00177778T³ + 0.0952381T² - 2.03175T + 24.2381', center=True, indent=False)
add_para(doc, '其中T表示贮藏温度，单位为℃；L₀表示温度基础货架期，单位为天。硬度修正通过分段线性插值计算，校准系数在参考点之间连续变化。')
add_heading(doc, '6.8 基准数据说明', 2)
add_table(doc, ['贮藏温度', '手持硬度', '基准货架期'], [
    ['0℃', '70', '37天'],
    ['5℃', '70', '25天'],
    ['25℃', '60', '2天'],
], [4.0, 4.0, 6.3])
add_para(doc, '三个基准组合在模型校准后能够精确命中。非基准温度和硬度由连续函数计算，使相邻输入条件之间保持平滑变化。')

# 7
new_chapter(doc, '7 数据分析')
add_heading(doc, '7.1 数据分析区域', 2)
add_para(doc, '数据分析与历史记录区域位于预测工作台下部。完成预测后该区域自动更新，展示当前账号最近100条记录的统计结果和历史数据。Excel导出包含当前账号的全部历史记录。')
add_figure(doc, '图7-1_数据分析与历史记录区域.png', '图7-1 数据分析与历史记录区域', max_width_cm=14.6, max_height_cm=5.0)
add_heading(doc, '7.2 预测次数统计', 2)
add_para(doc, '“预测次数”显示当前分析范围内的历史记录总数。页面分析范围最多为最近100条，超过100条的数据仍保存在数据库中，并可通过Excel导出。')
add_heading(doc, '7.3 平均货架期统计', 2)
add_para(doc, '“平均货架期”根据当前分析范围内每条记录的预计货架期计算算术平均值，结果保留一位小数，单位为天。无历史记录时显示0。')
add_heading(doc, '7.4 低风险记录统计', 2)
add_para(doc, '“低风险”显示当前分析范围内风险等级为低的记录数量。该指标用于查看适宜冷链贮藏条件所占数量。')
add_heading(doc, '7.5 中风险记录统计', 2)
add_para(doc, '“中风险”显示当前分析范围内风险等级为中的记录数量。中风险记录应优先安排销售，并减少温度波动。')
add_heading(doc, '7.6 高风险记录统计', 2)
add_para(doc, '“高风险”显示当前分析范围内风险等级为高的记录数量。高风险记录应缩短巡检间隔，并及时安排销售、加工或剔除异常果实。')
add_figure(doc, '图7-2_统计指标.png', '图7-2 预测统计指标', max_width_cm=14.6, max_height_cm=4.0)
add_heading(doc, '7.7 货架期趋势图', 2)
add_para(doc, '货架期变化趋势图按照预测时间展示最近20次预计货架期。横轴表示预测日期，纵轴表示预计货架期天数。趋势图可用于比较连续预测结果的变化方向。记录不足20条时显示现有全部记录。')
add_heading(doc, '7.8 风险等级分布图', 2)
add_para(doc, '风险分布图统计最近100条记录中低风险、中风险和高风险的数量占比，并以百分比和横向进度条表示。无历史记录时，各项比例均显示0%。')
add_figure(doc, '图7-3_趋势图与风险分布图.png', '图7-3 货架期趋势图与风险分布图', max_width_cm=14.6, max_height_cm=9.5)
add_heading(doc, '7.9 图表数据说明', 2)
add_para(doc, '图表数据与历史记录来源一致。每次预测保存成功后，页面重新读取当前账号的历史数据并重绘图表。浏览器窗口尺寸发生变化时，趋势图按照可用宽度重新绘制。')

# 8
new_chapter(doc, '8 历史记录管理')
add_heading(doc, '8.1 历史记录列表', 2)
add_para(doc, '历史记录按照预测时间倒序排列，最新预测显示在列表顶部。页面最多展示最近100条记录。每条记录包含预测输入、预测结果、风险、置信度和方法信息。')
add_figure(doc, '图8-1_预测历史标题区域.png', '图8-1 预测历史区域', max_width_cm=14.6, max_height_cm=3.5)
add_heading(doc, '8.2 历史记录字段说明', 2)
add_table(doc, ['字段', '说明'], [
    ['预测时间', '预测记录保存时的日期和时间。'],
    ['温度', '提交预测时选择的贮藏温度，单位为℃。'],
    ['手持硬度', '提交预测时输入的硬度计读数。'],
    ['映射硬度', '系统换算后的硬度值，单位为g·mm⁻²。'],
    ['预计货架期', '预测模型计算得到的中心结果，单位为天。'],
    ['预测区间', '预计货架期的参考下限和上限。'],
    ['风险', '低、中或高风险等级。'],
    ['置信度', '以百分比显示的参考置信度。'],
    ['方法', '本条记录采用的预测方法。'],
], [4.0, 10.3])
add_heading(doc, '8.3 查看历史预测数据', 2)
add_para(doc, '向下滚动页面进入预测历史区域，可横向查看全部字段。新增预测保存成功后，最新记录自动插入列表顶部。历史记录与登录账号绑定，切换账号后显示对应账号的数据。')
add_figure(doc, '图8-2_历史记录列表.png', '图8-2 历史记录列表', max_width_cm=14.6, max_height_cm=8.5)
add_heading(doc, '8.4 刷新历史记录', 2)
add_para(doc, '预测成功后系统自动刷新历史记录。需要重新读取数据库数据时，可刷新浏览器页面。刷新页面不会重复保存预测记录，也不会清除已经保存的数据。')
add_heading(doc, '8.5 导出Excel文件', 2)
add_steps(doc, [
    '登录需要导出数据的用户账号。',
    '进入“数据分析与历史记录”区域。',
    '选择右上角的“导出Excel”按钮。',
    '浏览器开始下载.xlsx文件。',
    '在浏览器下载目录中打开或保存该文件。',
])
add_heading(doc, '8.6 Excel文件字段说明', 2)
add_para(doc, '导出的工作簿名为预测历史，第一行为表头。文件对表头进行固定、筛选和列宽设置，可使用Microsoft Excel或其他兼容软件打开。')
add_table(doc, ['序号', 'Excel字段'], [
    ['1', '序号'], ['2', '预测时间（UTC）'], ['3', '用户'], ['4', '贮藏温度（℃）'],
    ['5', '手持硬度'], ['6', '映射硬度（g·mm⁻²）'], ['7', '预计货架期（天）'],
    ['8', '区间下限（天）'], ['9', '区间上限（天）'], ['10', '风险等级'],
    ['11', '置信度'], ['12', '预测方法'], ['13', '处理建议'],
], [2.4, 11.9])
add_heading(doc, '8.7 清除历史记录', 2)
add_steps(doc, [
    '进入“数据分析与历史记录”区域。',
    '选择右上角的“清除历史”按钮。',
    '阅读浏览器显示的确认信息。',
    '确认清除后，系统删除当前账号的全部预测历史。',
    '页面重新加载，统计指标、图表和列表恢复为空数据状态。',
])
add_heading(doc, '8.8 清除操作确认', 2)
add_para(doc, '清除操作执行前会显示确认提示。选择取消时不删除任何数据；选择确认后执行删除。历史清除后无法通过页面恢复，需要保留的数据应先导出Excel文件。')
add_figure(doc, '图8-3_导出与清除功能.png', '图8-3 Excel导出与历史清除入口', max_width_cm=14.6, max_height_cm=4.0)
add_heading(doc, '8.9 用户历史数据隔离', 2)
add_para(doc, '数据库中的每条预测记录均保存所属用户标识。历史查询、统计、Excel导出和清除操作均使用当前登录用户标识作为数据条件，从而避免不同账号之间相互查看或操作历史数据。')

# 9
new_chapter(doc, '9 输入校验与异常处理')
add_heading(doc, '9.1 温度输入范围', 2)
add_para(doc, '浏览器页面提供0℃至25℃范围内的六个预设温度按钮。后端预测接口校验温度是否位于0℃至25℃范围内，超出范围的请求不执行预测。')
add_heading(doc, '9.2 硬度输入范围', 2)
add_para(doc, '手持硬度允许范围为50至90。低于50、高于90、为空或无法转换为数值的输入不符合预测要求，系统不保存对应记录。')
add_heading(doc, '9.3 必填项校验', 2)
add_para(doc, '温度、手持硬度和预测方法均为预测必需参数。页面通过输入控件和请求校验共同保证参数完整。提交按钮在计算期间暂时禁用，避免短时间内重复提交。')
add_heading(doc, '9.4 用户名校验', 2)
add_para(doc, '用户名长度为3至30个字符，只能包含中文、字母、数字、下划线或连字符。用户名在系统中保持唯一。')
add_heading(doc, '9.5 密码校验', 2)
add_para(doc, '注册密码长度不少于8个字符，两次密码输入必须一致。系统仅保存密码哈希，不直接保存密码明文。')
add_heading(doc, '9.6 登录状态失效', 2)
add_para(doc, '未登录访问预测工作台时，系统自动跳转至登录页面。登录会话失效后提交预测，页面返回登录流程。重新登录后可继续使用原账号历史数据。')
add_heading(doc, '9.7 预测请求异常', 2)
add_para(doc, '预测请求失败时，参数区域显示错误信息，提交按钮恢复可用状态。应检查输入范围、预测方法状态和服务运行状态后重新提交。失败请求不会增加历史记录。')
add_heading(doc, '9.8 数据导出异常', 2)
add_para(doc, 'Excel文件未下载时，应检查浏览器下载权限、默认下载目录和服务运行状态。无历史数据时，导出的工作簿仍包含规范表头，不包含预测数据行。')
add_heading(doc, '9.9 页面刷新与重新登录', 2)
add_para(doc, '页面显示异常时可使用浏览器刷新功能重新加载。刷新后系统重新读取数据库中的当前用户历史记录。账号状态失效时应返回登录页面重新验证。')
add_table(doc, ['异常现象', '可能原因', '处理方式'], [
    ['无法进入工作台', '未登录或会话失效', '重新登录账号。'],
    ['预测按钮无响应', '参数不完整或服务未运行', '检查输入并确认服务状态。'],
    ['预测结果未保存', '请求失败或登录状态失效', '查看错误提示后重新提交。'],
    ['图表无数据', '当前账号没有历史记录', '完成至少一次预测。'],
    ['Excel未下载', '浏览器限制下载或服务异常', '允许下载并重新执行导出。'],
    ['历史数据为空', '登录了其他账号或执行过清除', '确认当前用户名和操作记录。'],
], [3.6, 5.0, 5.7])

# 10
new_chapter(doc, '10 数据安全与使用注意事项')
add_heading(doc, '10.1 用户账号安全', 2)
add_para(doc, '不同使用人员应使用独立账号。公共设备操作完成后应退出登录，不应将账号密码交由无关人员使用。')
add_heading(doc, '10.2 密码安全', 2)
add_para(doc, '密码长度应满足系统要求，并避免使用容易猜测的连续数字、姓名或简单重复字符。系统使用密码哈希完成验证，数据库中不保存可直接读取的密码明文。')
add_heading(doc, '10.3 用户数据隔离', 2)
add_para(doc, '预测、查询、统计、导出和清除均以当前登录用户为数据范围。账号切换后，页面显示新账号对应的数据。')
add_heading(doc, '10.4 请求安全校验', 2)
add_para(doc, '注册、登录、退出、预测和清除等页面提交操作使用安全令牌进行来源校验。页面失效或令牌不正确时，系统拒绝相关请求。')
add_heading(doc, '10.5 历史数据管理', 2)
add_para(doc, '重要预测历史应定期导出Excel并按照业务要求归档。数据库文件应随系统数据备份计划进行备份，避免设备故障导致历史记录丢失。')
add_heading(doc, '10.6 数据清除注意事项', 2)
add_para(doc, '“清除历史”将删除当前账号的全部预测记录，并同步清空统计和图表数据。该操作不可通过系统界面撤销。执行清除前应确认已完成必要的数据导出。')
add_heading(doc, '10.7 预测结果使用说明', 2)
add_para(doc, '预测结果以输入参数和系统模型为基础，用于贮藏管理和销售安排参考。实际货架期还可能受到采收成熟度、包装、运输振动、环境湿度、病害和个体差异影响。使用人员应结合现场检查、硬度复测和质量管理要求作出最终处理决定。')

# 11
new_chapter(doc, '11 常见问题')
faq = [
    ('11.1 无法进入预测工作台', '未登录或登录状态已失效时，系统会进入登录页面。输入正确的用户名和密码完成登录后即可访问工作台。'),
    ('11.2 注册信息无法提交', '检查用户名长度和字符范围，确认密码不少于8个字符，并确认两次密码输入一致。用户名已存在时应更换用户名或使用已有账号登录。'),
    ('11.3 预测结果未显示', '检查温度是否已经选择、硬度是否在50至90之间，并确认预测方法处于可用状态。服务停止时需重新启动软件。'),
    ('11.4 历史记录未更新', '等待预测状态显示已保存后查看列表。页面未自动刷新时，可刷新浏览器重新读取历史数据。'),
    ('11.5 Excel文件无法下载', '检查浏览器是否阻止下载，确认下载目录具有写入权限，并重新选择“导出Excel”。'),
    ('11.6 清除后能否恢复数据', '清除操作会从数据库删除当前账号的历史记录，系统页面不提供恢复功能。重要数据应在清除前导出。'),
    ('11.7 如何退出当前账号', '选择页面右上角的“退出”按钮。退出后系统返回登录页面，当前浏览器会话不再保持登录状态。'),
]
for title, text in faq:
    add_heading(doc, title, 2)
    add_para(doc, text)

# Appendix A
new_chapter(doc, '附录A 软件功能结构')
add_heading(doc, 'A.1 功能模块组成', 2)
add_table(doc, ['一级模块', '二级功能'], [
    ['用户管理', '用户注册、用户登录、用户退出、会话验证'],
    ['预测管理', '温度选择、硬度输入、硬度映射、方法选择、结果计算'],
    ['结果展示', '预计货架期、预测区间、风险等级、置信度、处理建议'],
    ['数据分析', '预测次数、平均货架期、风险数量、趋势图、风险分布'],
    ['历史管理', '历史列表、Excel导出、历史清除、用户数据隔离'],
    ['安全控制', '密码哈希、登录访问控制、安全令牌校验'],
], [4.0, 10.3])
add_heading(doc, 'A.2 用户操作流程', 2)
add_steps(doc, [
    '打开系统访问地址。',
    '注册新账号或登录已有账号。',
    '进入预测工作台。',
    '设置温度和硬度参数。',
    '提交预测并查看结果。',
    '查看统计图表和历史记录。',
    '导出或清除历史数据。',
    '退出系统。',
])
add_heading(doc, 'A.3 预测业务流程', 2)
add_steps(doc, [
    '接收温度、手持硬度和预测方法。',
    '校验输入范围和方法状态。',
    '将手持硬度映射为标准硬度值。',
    '计算温度基础货架期。',
    '应用硬度修正和平滑校准。',
    '计算预计货架期和预测区间。',
    '确定风险等级、置信度和处理建议。',
    '保存当前用户预测历史。',
    '刷新统计指标、图表和历史列表。',
])

# Appendix B
new_chapter(doc, '附录B 数据字段说明')
add_heading(doc, 'B.1 预测输入字段', 2)
add_table(doc, ['字段', '类型', '范围或取值', '说明'], [
    ['贮藏温度', '数值', '0℃至25℃', '页面提供六个预设温度。'],
    ['手持硬度', '数值', '50至90', '手持式硬度计读数。'],
    ['硬度单位', '文本', 'handheld', '表示手持式硬度读数。'],
    ['预测方法', '文本', '多项式回归等', '页面仅允许选择当前可用方法。'],
], [3.0, 2.3, 4.0, 5.0])
add_heading(doc, 'B.2 预测结果字段', 2)
add_table(doc, ['字段', '说明'], [
    ['预计货架期', '模型计算得到的中心预测结果，单位为天。'],
    ['区间下限', '预测区间的最小参考天数。'],
    ['区间上限', '预测区间的最大参考天数。'],
    ['风险等级', '低、中或高。'],
    ['置信度', '0至1之间的数值，页面转换为百分比显示。'],
    ['处理建议', '与风险等级对应的销售或贮藏建议。'],
    ['操作建议', '包装、温度控制、复测和巡检建议。'],
], [4.0, 10.3])
add_heading(doc, 'B.3 历史记录字段', 2)
add_para(doc, '历史记录除保存预测输入和结果字段外，还保存记录编号、所属用户、预测方法和创建时间。历史记录按照创建时间和记录编号倒序显示。')
add_heading(doc, 'B.4 Excel导出字段', 2)
add_para(doc, 'Excel导出字段包括序号、预测时间、用户、贮藏温度、手持硬度、映射硬度、预计货架期、区间下限、区间上限、风险等级、置信度、预测方法和处理建议。')

# Appendix C
new_chapter(doc, '附录C 软件界面图索引')
add_table(doc, ['图号', '图名', '所在章节'], [
    ['图2-1', '系统主界面全景', '第2章'],
    ['图4-1', '用户登录界面', '第4章'],
    ['图4-2', '用户注册界面', '第4章'],
    ['图5-1', '预测工作台', '第5章'],
    ['图5-2', '预测参数与结果区域', '第5章'],
    ['图6-1', '预测结果界面', '第6章'],
    ['图7-1', '数据分析与历史记录区域', '第7章'],
    ['图7-2', '预测统计指标', '第7章'],
    ['图7-3', '货架期趋势图与风险分布图', '第7章'],
    ['图8-1', '预测历史区域', '第8章'],
    ['图8-2', '历史记录列表', '第8章'],
    ['图8-3', 'Excel导出与历史清除入口', '第8章'],
], [2.6, 7.7, 4.0])
add_para(doc, '以上界面图均为软件实际运行界面，界面内容与本说明书对应功能一致。')

# Enforce black text throughout and no accidental hyperlink color.
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        run.font.color.rgb = BLACK
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = BLACK

# Set update fields on open.
settings = doc.settings.element
update_fields = settings.find(qn('w:updateFields'))
if update_fields is None:
    update_fields = OxmlElement('w:updateFields')
    settings.append(update_fields)
update_fields.set(qn('w:val'), 'true')

doc.save(OUT_FILE)
print(OUT_FILE)
print('size', OUT_FILE.stat().st_size)
print('paragraphs', len(doc.paragraphs), 'tables', len(doc.tables), 'sections', len(doc.sections))
