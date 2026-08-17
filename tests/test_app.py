from __future__ import annotations

import io
import re
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from app import AUTH_ENABLED, HISTORY_ENABLED, create_app


AUTH_HISTORY_ENABLED = AUTH_ENABLED and HISTORY_ENABLED


class MeizaoAppTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_directory = tempfile.TemporaryDirectory()
        database_path = str(Path(self.temp_directory.name) / "test.sqlite3")
        self.app = create_app(
            {
                "TESTING": True,
                "DATABASE": database_path,
                "SECRET_KEY": "test-secret",
                "CSRF_ENABLED": False,
            }
        )
        self.client = self.app.test_client()

    def tearDown(self) -> None:
        self.temp_directory.cleanup()

    def register(self, username: str = "测试用户", password: str = "StrongPass123"):
        return self.client.post(
            "/register",
            data={
                "username": username,
                "password": password,
                "confirm_password": password,
            },
            follow_redirects=True,
        )

    @unittest.skipUnless(AUTH_HISTORY_ENABLED, "登录和历史记录功能已临时关闭")
    def test_authentication_prediction_history_and_excel_export(self) -> None:
        self.assertEqual(self.client.get("/").status_code, 302)
        response = self.register()
        self.assertEqual(response.status_code, 200)
        self.assertIn("预测工作台", response.get_data(as_text=True))

        prediction_response = self.client.post(
            "/predict",
            json={
                "storage_temperature_c": 5,
                "firmness": 70,
                "firmness_unit": "handheld",
                "prediction_method": "polynomial_regression",
            },
        )
        self.assertEqual(prediction_response.status_code, 200)
        prediction = prediction_response.get_json()
        self.assertEqual(prediction["history_id"], 1)

        history_response = self.client.get("/api/history")
        history = history_response.get_json()
        self.assertEqual(history["summary"]["total"], 1)
        self.assertEqual(
            history["records"][0]["estimated_shelf_life_days"],
            prediction["estimated_shelf_life_days"],
        )

        export_response = self.client.get("/history/export")
        self.assertEqual(export_response.status_code, 200)
        workbook = load_workbook(io.BytesIO(export_response.data))
        worksheet = workbook["预测历史"]
        self.assertEqual(worksheet.max_row, 2)
        self.assertEqual(worksheet["C2"].value, "测试用户")

    def test_index_allows_continuous_temperature_input(self) -> None:
        page = self.client.get("/")
        self.assertEqual(page.status_code, 200)
        html = page.get_data(as_text=True)
        self.assertIn('id="storageTemperature"', html)
        self.assertIn('min="0" max="25" step="any"', html)

        response = self.client.post(
            "/predict",
            json={
                "storage_temperature_c": 12.345,
                "firmness": 70,
                "firmness_unit": "handheld",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["storage_temperature_c"], 12.345)

    def test_reference_points_and_continuous_generalization(self) -> None:
        self.register()
        for temperature, hardness, expected_days in (
            (0, 70, 37.0),
            (5, 70, 25.0),
            (25, 60, 2.0),
        ):
            response = self.client.post(
                "/predict",
                json={
                    "storage_temperature_c": temperature,
                    "firmness": hardness,
                    "firmness_unit": "handheld",
                    "prediction_method": "polynomial_regression",
                },
            )
            self.assertEqual(response.status_code, 200)
            self.assertEqual(
                response.get_json()["estimated_shelf_life_days"], expected_days
            )

        interpolated = self.client.post(
            "/predict",
            json={
                "storage_temperature_c": 2.5,
                "firmness": 70,
                "firmness_unit": "handheld",
                "prediction_method": "polynomial_regression",
            },
        )
        self.assertEqual(interpolated.status_code, 200)
        interpolated_days = interpolated.get_json()["estimated_shelf_life_days"]
        self.assertGreater(interpolated_days, 25.0)
        self.assertLess(interpolated_days, 37.0)

        nearby_hardness_results = []
        for hardness in (60, 65, 70):
            response = self.client.post(
                "/predict",
                json={
                    "storage_temperature_c": 25,
                    "firmness": hardness,
                    "firmness_unit": "handheld",
                    "prediction_method": "polynomial_regression",
                },
            )
            self.assertEqual(response.status_code, 200)
            nearby_hardness_results.append(
                response.get_json()["estimated_shelf_life_days"]
            )
        self.assertLess(nearby_hardness_results[0], nearby_hardness_results[1])
        self.assertLess(nearby_hardness_results[1], nearby_hardness_results[2])

    @unittest.skipUnless(AUTH_HISTORY_ENABLED, "登录和历史记录功能已临时关闭")
    def test_each_user_only_sees_their_own_history(self) -> None:
        self.register("first_user")
        self.client.post(
            "/predict",
            json={
                "storage_temperature_c": 10,
                "firmness": 65,
                "firmness_unit": "handheld",
                "prediction_method": "polynomial_regression",
            },
        )
        self.client.post("/logout")
        self.register("second_user")
        history = self.client.get("/api/history").get_json()
        self.assertEqual(history["summary"]["total"], 0)
        self.assertEqual(history["records"], [])

    @unittest.skipUnless(AUTH_HISTORY_ENABLED, "登录和历史记录功能已临时关闭")
    def test_clear_history_only_removes_current_users_records(self) -> None:
        self.register("first_user")
        self.client.post(
            "/predict",
            json={
                "storage_temperature_c": 10,
                "firmness": 65,
                "firmness_unit": "handheld",
                "prediction_method": "polynomial_regression",
            },
        )
        self.client.post("/logout")

        self.register("second_user")
        self.client.post(
            "/predict",
            json={
                "storage_temperature_c": 5,
                "firmness": 70,
                "firmness_unit": "handheld",
                "prediction_method": "polynomial_regression",
            },
        )
        cleared = self.client.post("/history/clear", follow_redirects=True)
        self.assertEqual(cleared.status_code, 200)
        self.assertIn("已清除 1 条预测历史", cleared.get_data(as_text=True))
        self.assertEqual(self.client.get("/api/history").get_json()["summary"]["total"], 0)

        self.client.post("/logout")
        login = self.client.post(
            "/login",
            data={"username": "first_user", "password": "StrongPass123"},
            follow_redirects=True,
        )
        self.assertEqual(login.status_code, 200)
        self.assertEqual(self.client.get("/api/history").get_json()["summary"]["total"], 1)

    @unittest.skipUnless(AUTH_ENABLED, "登录注册功能已临时关闭")
    def test_csrf_protection_rejects_missing_token(self) -> None:
        csrf_app = create_app(
            {
                "TESTING": True,
                "DATABASE": str(Path(self.temp_directory.name) / "csrf.sqlite3"),
                "SECRET_KEY": "csrf-secret",
                "CSRF_ENABLED": True,
            }
        )
        client = csrf_app.test_client()
        page = client.get("/register")
        match = re.search(r'name="csrf_token" value="([^"]+)"', page.get_data(as_text=True))
        self.assertIsNotNone(match)

        rejected = client.post(
            "/register",
            data={
                "username": "csrf_user",
                "password": "StrongPass123",
                "confirm_password": "StrongPass123",
            },
        )
        self.assertEqual(rejected.status_code, 400)

        accepted = client.post(
            "/register",
            data={
                "csrf_token": match.group(1),
                "username": "csrf_user",
                "password": "StrongPass123",
                "confirm_password": "StrongPass123",
            },
        )
        self.assertEqual(accepted.status_code, 302)


if __name__ == "__main__":
    unittest.main()
