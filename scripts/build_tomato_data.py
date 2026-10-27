"""Build auditable tomato inputs from the supplied experimental workbook and raw readings."""
import argparse, csv, hashlib, json
from pathlib import Path
from statistics import mean
from openpyxl import load_workbook


def build(source):
    workbook=source/'小番茄货架期数据.xlsx'
    raw=source/'小番茄硬度数据.csv'
    w=load_workbook(workbook,data_only=True)
    ratios=[]
    for row in w['百分比换算组'].iter_rows(min_row=2,values_only=True):
        day=int(str(row[0]).rstrip('d'))
        for temp,value in zip(range(0,26,5),row[1:]):
            if value is not None:
                ratios.append(dict(temperature=temp,day=day,ratio=float(value)))
    repeats=[]
    for row in w['五个重复组'].iter_rows(min_row=2,values_only=True):
        day=int(str(row[0]).rstrip('d'))
        for temp,value in zip(range(0,26,5),row[1:]):
            if value is not None:
                counts=[int(v) for v in str(value).split('/')] if '/' in str(value) else [int(value)]*5
                assert len(counts)==5
                repeats.append(dict(temperature=temp,day=day,counts=counts))
    counts={(r['temperature'],r['day']):r['counts'] for r in repeats}
    for r in ratios:
        assert abs(sum(counts[r['temperature'],r['day']])/125-r['ratio'])<1e-10
    groups={}
    for row in csv.DictReader(raw.open(encoding='utf-8-sig')):
        for metric,col in [('g_per_mm','硬度（g/mm）'),('integer','硬度（整数）')]:
            if row[col]:groups.setdefault((int(row['温度']),int(row['时间']),metric),[]).append(float(row[col]))
    firmness=[dict(temperature=t,day=d,metric=m,n=len(v),mean=mean(v)) for (t,d,m),v in sorted(groups.items())]
    parameter_file=source/'拟合补全_模型估计/最终补全参数.json'
    params=json.loads(parameter_file.read_text(encoding='utf-8'))
    coefficients={m:params[col]['soft_coef'][2:] for m,col in [('g_per_mm','硬度（g/mm）'),('integer','硬度（整数）')]}
    data=dict(version='tomato-2026-10-27-v1',ratios=ratios,repeat_counts=repeats,firmness=firmness,scenario_coefficients=coefficients,
      source_sha256={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in (workbook,raw,parameter_file)},
      provenance={'ratio_sheet':'百分比换算组','repeat_sheet':'五个重复组','raw_firmness_readings':sum(len(v) for v in groups.values()),
        'ratio_interpretation':'比例指标的具体实验判定口径未确认；阈值为用户设定的探索性标准，不是食品安全界限。',
        'denominator_note':'已有比例与五组计数总和/125一致，仅用于校验；不推定遗漏比例值。',
        'missing_policy':'原工作簿空白保持缺失；不从重复组补填缺失比例；硬度均值仅使用原始非空读数。'})
    return data

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('source',type=Path);a=p.parse_args()
    out=Path(__file__).resolve().parents[1]/'data/tomato_observations.json'
    out.write_text(json.dumps(build(a.source),ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(out)
